from __future__ import annotations

import argparse
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "config"
OUTPUT = ROOT / "output"

def load_json(path: Path):
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)

def player_root(raw: dict) -> dict:
    return raw.get("player", raw)

def load_units(player: dict) -> list[dict]:
    units = player.get("units")
    if not isinstance(units, list):
        raise ValueError("player.json does not contain a units array")
    return units

def ability_pair(unit: dict):
    abilities = unit.get("abilities", [])
    if len(abilities) < 2:
        return None
    return abilities[0], abilities[1]

def add_table(ws, ref: str, name: str):
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tab)

def style_sheet(ws, widths: dict[str, int]):
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

def build(player_path: Path, output_path: Path):
    player = player_root(load_json(player_path))
    units = load_units(player)
    by_name = {u.get("name"): u for u in units}

    priorities = load_json(CONFIG / "character_priorities.json")
    targets = load_json(CONFIG / "ability_targets.json")
    watch = load_json(CONFIG / "equipment_watchlist.json")

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Legendary Watchlist")
    headers = ["Priority","Legendary Item","Character","Slot","Buy Tier","Character Value","Current Character Rank","Notes"]
    ws.append(headers)
    for row in watch:
        char = row["character"]
        u = by_name.get(char, {})
        cp = priorities.get(char, {})
        ws.append([
            row["priority"], row["item"], char, row["slot"], row["tier"],
            cp.get("priority",""), u.get("rank",""), cp.get("note","")
        ])
    if ws.max_row > 1:
        add_table(ws, f"A1:H{ws.max_row}", "LegendaryWatchlist")
    style_sheet(ws, {"A":9,"B":28,"C":20,"D":16,"E":14,"F":16,"G":20,"H":60})

    wa = wb.create_sheet("Ability Priorities")
    wa.append(["Character","Faction","Rank","Active ID","Active Lv","Passive ID","Passive Lv",
               "Uncommon Priority","Long-Term Focus","Active Target","Passive Target","Recommendation Basis"])
    queue = []
    for u in units:
        pair = ability_pair(u)
        if not pair:
            continue
        active, passive = pair
        name = u.get("name","")
        cp = priorities.get(name,{})
        target = targets.get(name)
        under_active = active.get("level",0) < 17
        under_passive = passive.get("level",0) < 17

        if under_active or under_passive:
            uncommon_priority = "HIGH" if cp.get("priority",0) >= 80 or target else "MEDIUM"
        else:
            uncommon_priority = "DONE / ABOVE 17"

        if target:
            active_target = target["active"]
            passive_target = target["passive"]
            focus = target["focus"]
            basis = target["confidence"]
        else:
            active_target = "17" if under_active else "Current+"
            passive_target = "17" if under_passive else "Current+"
            focus = "Baseline / review later"
            basis = "User level-17 baseline; no researched character-specific target stored"

        wa.append([name,u.get("faction",""),u.get("rank",""),active.get("id",""),active.get("level",0),
                   passive.get("id",""),passive.get("level",0),uncommon_priority,focus,
                   active_target,passive_target,basis])

        if under_active or under_passive:
            queue.append([name,u.get("faction",""),active.get("id",""),active.get("level",0),
                          passive.get("id",""),passive.get("level",0),uncommon_priority,
                          "YES" if under_active else "No","YES" if under_passive else "No",focus])

    if wa.max_row > 1:
        add_table(wa, f"A1:L{wa.max_row}", "AbilityPriorities")
    style_sheet(wa, {"A":20,"B":20,"C":10,"D":28,"E":10,"F":28,"G":10,"H":18,"I":20,"J":18,"K":18,"L":48})

    wq = wb.create_sheet("Uncommon Badge Queue")
    wq.append(["Character","Faction","Active ID","Active Lv","Passive ID","Passive Lv",
               "Priority","Active to 17?","Passive to 17?","Long-Term Focus"])
    for row in sorted(queue, key=lambda x: (0 if x[6]=="HIGH" else 1, x[0])):
        wq.append(row)
    if wq.max_row > 1:
        add_table(wq, f"A1:J{wq.max_row}", "UncommonQueue")
    style_sheet(wq, {"A":20,"B":20,"C":28,"D":10,"E":28,"F":10,"G":14,"H":16,"I":16,"J":22})

    wr = wb.create_sheet("Read Me")
    wr.append(["Tacticus Upgrade Helper",""])
    wr.append(["Roster source", str(player_path)])
    wr.append(["Ability policy","Raise unlocked abilities to 9 immediately; current project is useful abilities through Uncommon to level 17."])
    wr.append(["Recommendation data","Stored separately under config/ so it can be improved without modifying player exports."])
    wr.append(["Privacy","player.json is gitignored and should not be committed."])
    style_sheet(wr, {"A":28,"B":90})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Wrote {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Generate account-specific Tacticus upgrade reports.")
    parser.add_argument("player_json", type=Path)
    parser.add_argument("-o","--output", type=Path, default=OUTPUT / "Tacticus_Upgrade_Priorities.xlsx")
    args = parser.parse_args()
    build(args.player_json, args.output)

if __name__ == "__main__":
    main()
